"""Build docs/HealthPro.pptx. Run from the HealthPro directory."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = Path(__file__).resolve().parent / "HealthPro.pptx"

GREEN = RGBColor(0x0F, 0x3D, 0x2E)
MID = RGBColor(0x1F, 0x6F, 0x54)
SAGE = RGBColor(0xE7, 0xF3, 0xEC)
CREAM = RGBColor(0xF7, 0xF4, 0xEA)
INK = RGBColor(0x14, 0x21, 0x1C)
MUTED = RGBColor(0x5B, 0x6B, 0x63)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
W, H = Inches(13.333), Inches(7.5)


def set_run(run, text, size, color, bold=False, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = font


def add_text(slide, text, x, y, w, h, size, color, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    set_run(p.add_run(), text, size, color, bold)
    return box


def rect(slide, x, y, w, h, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def footer(slide, number, total):
    rect(slide, 0, Inches(7.15), W, Inches(0.35), GREEN)
    add_text(slide, "HealthPro  ·  Personal nutrition tracker  ·  Educational, not medical advice",
             Inches(0.45), Inches(7.18), Inches(10), Inches(0.28), 11, WHITE)
    add_text(slide, f"{number}  /  {total}", Inches(11.4), Inches(7.18), Inches(1.5), Inches(0.28),
             11, WHITE, align=PP_ALIGN.RIGHT)


def header(slide, kicker, title):
    rect(slide, 0, 0, Inches(0.18), H, MID)
    add_text(slide, kicker.upper(), Inches(0.55), Inches(0.28), Inches(12), Inches(0.28), 12, MID, True)
    add_text(slide, title, Inches(0.55), Inches(0.52), Inches(12.2), Inches(0.7), 30, GREEN, True)


def bullets(slide, items, x, y, w, h, size=18):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(10)
        p.level = 0
        set_run(p.add_run(), item, size, INK)
    return box


def card(slide, x, y, w, h, title, body, fill=SAGE):
    shape = rect(slide, x, y, w, h, fill)
    add_text(slide, title, x + Inches(0.2), y + Inches(0.12), w - Inches(0.35), Inches(0.36), 16, GREEN, True)
    add_text(slide, body, x + Inches(0.2), y + Inches(0.48), w - Inches(0.35), h - Inches(0.58), 14, INK)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    prs.core_properties.title = "HealthPro"
    prs.core_properties.subject = "Personal nutrition tracker"
    prs.core_properties.category = "Project presentation"
    blank = prs.slide_layouts[6]
    slides = []

    def new():
        slide = prs.slides.add_slide(blank)
        slides.append(slide)
        return slide

    # 1 Title
    s = new()
    rect(s, 0, 0, W, H, GREEN)
    rect(s, 0, 0, Inches(0.28), H, MID)
    add_text(s, "PERSONAL PROJECT", Inches(0.8), Inches(1.7), Inches(10), Inches(0.3), 14, SAGE, True)
    add_text(s, "HealthPro", Inches(0.75), Inches(2.05), Inches(11), Inches(1.1), 60, WHITE, True)
    add_text(s, "A personal food log that prices meals from food tables,\nthen compares the day with the person who ate them.",
             Inches(0.8), Inches(3.3), Inches(10), Inches(1.1), 22, SAGE)
    add_text(s, "React  ·  FastAPI  ·  Local JSON  ·  Version 0.1.0",
             Inches(0.8), Inches(5.5), Inches(10), Inches(0.4), 16, WHITE)

    # 2 Problem
    s = new()
    header(s, "Why it exists", "A calorie app is not enough")
    bullets(s, [
        "Indian home foods are often missing, or returned under a different name.",
        "A single food lookup does not say whether the day was short of iron, sodium, or calories for this person.",
        "Hypertension, diabetes, pregnancy, anemia, kidney disease, and thyroid disease change the target.",
        "A chatbot can describe a meal. It should not be the source of the calorie number.",
        "Recipes and diet plans are useful only when they use the same limits as the log.",
    ], Inches(0.6), Inches(1.55), Inches(12), Inches(5), 22)

    # 3 What it is
    s = new()
    header(s, "The product", "One local application, six pages")
    pages = [
        ("01", "Food lookup", "Name and quantity to energy and micronutrients."),
        ("02", "People", "Body, conditions, cuisine, and a saved diet plan."),
        ("03", "Daily log", "Foods, exercise, and the day’s balance."),
        ("04", "Trends", "Saved days, gaps, and a rough weight note."),
        ("05", "Recipe generator", "A recipe priced from its ingredients."),
        ("06", "Saved recipes", "Your own method, repriced on every save."),
    ]
    for i, (num, title, body) in enumerate(pages):
        col, row = i % 3, i // 3
        x = Inches(0.5 + col * 4.2)
        y = Inches(1.6 + row * 2.5)
        rect(s, x, y, Inches(3.95), Inches(2.25), SAGE)
        add_text(s, num, x + Inches(0.2), y + Inches(0.15), Inches(1.2), Inches(0.4), 18, MID, True)
        add_text(s, title, x + Inches(0.2), y + Inches(0.6), Inches(3.5), Inches(0.45), 20, GREEN, True)
        add_text(s, body, x + Inches(0.2), y + Inches(1.15), Inches(3.5), Inches(0.85), 15, INK)

    # 4 Rule
    s = new()
    header(s, "The rule", "Numbers from tables. Language from the model.")
    rows = [
        ("Starter catalog and optional IFCT import", "Measured-style values you can point at"),
        ("USDA FoodData Central", "Used when the local catalog has no exact food"),
        ("Gemini estimate", "Only if you ask, then cached and labelled ai_estimate"),
        ("BMR, TDEE, targets, diet chart", "Calculated in the backend. No model call"),
        ("Recipe wording", "Gemini writes steps. Ingredient weights are priced locally"),
    ]
    for i, (left, right) in enumerate(rows):
        y = Inches(1.55 + i * 1.0)
        fill = WHITE if i % 2 else SAGE
        rect(s, Inches(0.55), y, Inches(12.2), Inches(0.88), fill)
        add_text(s, left, Inches(0.8), y + Inches(0.22), Inches(5.6), Inches(0.5), 18, GREEN, True)
        add_text(s, right, Inches(6.6), y + Inches(0.22), Inches(5.8), Inches(0.5), 18, INK)

    # 5 Person
    s = new()
    header(s, "Personal targets", "The same person drives the log and the plan")
    card(s, Inches(0.5), Inches(1.6), Inches(4.0), Inches(4.9), "Profile",
         "Age, sex, height, weight, and activity produce BMR and TDEE.\n\nNationality and ethnicity select a customary cuisine. They are not treated as a biological requirement.")
    card(s, Inches(4.7), Inches(1.6), Inches(4.0), Inches(4.9), "Conditions",
         "Hypertension tightens sodium.\nDiabetes tightens carbohydrate and raises fiber.\nKidney, pregnancy, anemia, and thyroid each change a small set of targets.",
         CREAM)
    card(s, Inches(8.9), Inches(1.6), Inches(3.9), Inches(4.9), "What you see",
         "Intake versus estimated need.\nEach nutrient against the personal target.\nA note when the log is probably incomplete, especially for salt.")

    # 6 Diet chart
    s = new()
    header(s, "Diet plan", "Instructions, plus a day you can cook")
    add_text(s, "Both parts are saved on the person. Click Diet plan again after nationality, conditions, or body size change.",
             Inches(0.55), Inches(1.4), Inches(12), Inches(0.4), 16, MUTED)
    slots = [
        ("06:30", "On waking", "Water"),
        ("08:00", "Breakfast", "2 idli and sambar\nor dosa, or upma"),
        ("11:00", "Mid-morning", "Fruit, buttermilk,\nor a few nuts"),
        ("13:30", "Lunch", "Rice, sambar,\nvegetable, curd"),
        ("16:30", "Evening", "Unsweetened tea\nand roasted chana"),
        ("19:30", "Dinner", "Less rice or millet,\ndal, vegetables"),
    ]
    for i, (time, meal, body) in enumerate(slots):
        x = Inches(0.4 + i * 2.15)
        rect(s, x, Inches(2.1), Inches(2.02), Inches(3.5), SAGE if i % 2 == 0 else WHITE)
        add_text(s, time, x + Inches(0.12), Inches(2.25), Inches(1.8), Inches(0.4), 18, MID, True)
        add_text(s, meal, x + Inches(0.12), Inches(2.7), Inches(1.8), Inches(0.7), 16, GREEN, True)
        add_text(s, body, x + Inches(0.12), Inches(3.5), Inches(1.8), Inches(1.8), 14, INK)
    add_text(s, "Example for a South Indian pattern. Each main meal offers two or three choices. Times can shift with the person’s day.",
             Inches(0.55), Inches(5.8), Inches(12), Inches(0.8), 15, MUTED)

    # 7 Recipes
    s = new()
    header(s, "Recipes", "Two kinds of advice, stored differently")
    card(s, Inches(0.5), Inches(1.6), Inches(6.0), Inches(4.8), "Shown for the selected person",
         "Good, limit, or skip.\n\nSays how much of this dish fits their day.\n\nStays on the generator page.\nIt is not written into the saved recipe, because the next person is different.")
    card(s, Inches(6.8), Inches(1.6), Inches(6.0), Inches(4.8), "Saved with the recipe",
         "People with hypertension, diabetes, kidney disease, pregnancy, anemia, or thyroid disease.\n\nAvoid, or a stated amount.\n\nRecomputed every time the recipe is saved.",
         CREAM)

    # 8 Architecture
    s = new()
    header(s, "Architecture", "Browser, API, files, two optional services")
    boxes = [
        (0.5, "React\nport 5173", SAGE),
        (3.7, "FastAPI\nport 8000", SAGE),
        (6.9, "JSON files\nbackend/data", CREAM),
        (10.1, "USDA\nGemini", WHITE),
    ]
    for x, label, fill in boxes:
        rect(s, Inches(x), Inches(2.2), Inches(2.7), Inches(1.6), fill)
        add_text(s, label, Inches(x + 0.15), Inches(2.5), Inches(2.4), Inches(1.1), 18, GREEN, True, PP_ALIGN.CENTER)
    add_text(s, "→", Inches(3.2), Inches(2.6), Inches(0.5), Inches(0.6), 28, MID, True)
    add_text(s, "→", Inches(6.4), Inches(2.6), Inches(0.5), Inches(0.6), 28, MID, True)
    add_text(s, "→", Inches(9.6), Inches(2.6), Inches(0.5), Inches(0.6), 28, MID, True)
    bullets(s, [
        "The browser never sees API keys. CORS allows only the local Vite origins.",
        "People, day logs, recipes, USDA responses, and AI estimates are written through to disk.",
        "USDA and Gemini are optional. Lookup, logs, trends, manual recipes, and diet plans work without them.",
    ], Inches(0.6), Inches(4.3), Inches(12), Inches(2.2), 18)

    # 9 Data
    s = new()
    header(s, "Your data", "Backup is a folder copy")
    files = [
        ("people.json", "Profiles and diet plans"),
        ("logs/<person>/<date>.json", "One saved day"),
        ("saved_recipes.json", "Recipes and condition notes"),
        ("usda_cache.json", "USDA responses"),
        ("ai_food_cache.json", "Labelled estimates"),
        ("ifct_import.json", "Optional official rows you add"),
    ]
    for i, (name, desc) in enumerate(files):
        col, row = i % 2, i // 2
        x = Inches(0.55 + col * 6.4)
        y = Inches(1.6 + row * 1.6)
        rect(s, x, y, Inches(6.1), Inches(1.4), SAGE)
        add_text(s, name, x + Inches(0.25), y + Inches(0.25), Inches(5.6), Inches(0.4), 18, GREEN, True)
        add_text(s, desc, x + Inches(0.25), y + Inches(0.75), Inches(5.6), Inches(0.4), 16, INK)

    # 10 Run
    s = new()
    header(s, "Run it", "Two commands, then the browser")
    rect(s, Inches(0.55), Inches(1.7), Inches(12.2), Inches(2.2), GREEN)
    add_text(s, "scripts\\start.cmd", Inches(0.85), Inches(1.95), Inches(11), Inches(0.55), 28, WHITE, True)
    add_text(s, "Starts the API and the UI in the background, waits until both answer,\nand prints  http://127.0.0.1:5173",
             Inches(0.85), Inches(2.65), Inches(11), Inches(0.9), 18, SAGE)
    card(s, Inches(0.55), Inches(4.2), Inches(6.0), Inches(2.3), "Stop",
         "scripts\\stop.cmd releases ports 8000 and 5173.")
    card(s, Inches(6.75), Inches(4.2), Inches(6.0), Inches(2.3), "Keys",
         "GEMINI_API_KEY and USDA_API_KEY live in .env. The health check reports only whether they are set.",
         CREAM)

    # 11 Limits
    s = new()
    header(s, "Limits", "What a careful reading should assume")
    bullets(s, [
        "The starter catalog is representative. It is not a copy of the ICMR–NIN tables. Those can be imported.",
        "Cups, katoris, and pieces are fixed gram guesses. Cooking loss is not modelled.",
        "The diet chart is a customary pattern with personal caps. It is not a prescription.",
        "Gemini’s free tier is small. When it is spent, generation stops and the rest of the app continues.",
        "The server is local and has no login. The machine user is the trust boundary.",
    ], Inches(0.6), Inches(1.6), Inches(12), Inches(4.8), 22)

    # 12 Close
    s = new()
    rect(s, 0, 0, W, H, GREEN)
    rect(s, 0, 0, Inches(0.28), H, MID)
    add_text(s, "HealthPro", Inches(0.8), Inches(1.8), Inches(11), Inches(0.9), 48, WHITE, True)
    add_text(s, "Track the day. Price the food. Keep the person in the calculation.",
             Inches(0.8), Inches(2.9), Inches(11), Inches(0.8), 22, SAGE)
    add_text(s, "Documents in  docs\\\nProposal, synopsis, SRS, HLD, LLD, final report, user manual, test plan.",
             Inches(0.8), Inches(4.2), Inches(11), Inches(1.2), 18, WHITE)
    add_text(s, "Educational personal tracker. Not medical advice.",
             Inches(0.8), Inches(6.2), Inches(11), Inches(0.4), 16, SAGE)

    total = len(slides)
    for i, slide in enumerate(slides, start=1):
        if i in (1, total):
            continue
        footer(slide, i, total)

    # Remove empty runs created by add_run on empty paragraphs? set_run sets text. OK.
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
